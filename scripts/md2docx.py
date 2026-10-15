#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown → Word (.docx) 导出工具
================================
把 notes.md 精读笔记转成 Word 文档。支持 notes.md 用到的语法子集：
YAML frontmatter（跳过）、#~#### 标题、段落（**粗** / *斜* / `代码`）、
无序/有序列表（含两级缩进）、GFM 表格、```代码块```、> 引用、--- 分割线、
![](图片)（相对路径基于 md 文件所在目录）、$$公式$$（保留 LaTeX 原文居中显示）。

用法:
  python md2docx.py notes.md [-o notes.docx]
依赖: python-docx (pip install python-docx)
"""

import argparse
import re
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    print("⚠ 缺少 python-docx 依赖，请运行: pip install python-docx")
    sys.exit(1)

INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`)")


def set_cjk(style_or_run, east: str, west: str = None):
    """给 run/样式同时设置中西文字体"""
    font = style_or_run.font if hasattr(style_or_run, "font") else style_or_run
    font.name = west or east
    rpr = font.element.get_or_add_rPr() if hasattr(font, "element") else None
    if rpr is not None:
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), east)


def add_runs(par, text: str, mono_ok: bool = True):
    """按 **粗** / *斜* / `code` 拆分加 run；其余原样"""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            r = par.add_run(part[2:-2]); r.bold = True; set_cjk(r, "宋体", "Times New Roman")
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"
            r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True; set_cjk(r, "宋体", "Times New Roman")
        else:
            r = par.add_run(part); set_cjk(r, "宋体", "Times New Roman")


def add_table(doc, rows):
    """rows: list[list[str]]（已去掉分隔行）"""
    if not rows:
        return
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = t.cell(ri, ci)
            cell.text = ""
            p = cell.paragraphs[0]
            add_runs(p, cell_text)
            for r in p.runs:
                r.font.size = Pt(10)
                if ri == 0:
                    r.bold = True
    doc.add_paragraph()


def split_table_row(line: str):
    return [c.strip() for c in line.strip().strip("|").split("|")]


def is_table_sep(line: str) -> bool:
    return bool(re.match(r"^\s*\|?[\s:|-]+\|?\s*$", line)) and "-" in line


def convert(md_path: Path, out_path: Path):
    doc = Document()
    # 基础样式
    normal = doc.styles["Normal"]
    normal.font.size = Pt(11)
    set_cjk(normal, "宋体", "Times New Roman")
    for name, size in (("Heading 1", 18), ("Heading 2", 15), ("Heading 3", 13), ("Heading 4", 12)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)
        set_cjk(st, "黑体", "Arial")

    lines = md_path.read_text(encoding="utf-8").splitlines()
    base_dir = md_path.parent

    # 跳过 frontmatter
    if lines and lines[0].strip() == "---":
        try:
            end = next(i for i in range(1, len(lines)) if lines[i].strip() == "---")
            lines = lines[end + 1:]
        except StopIteration:
            pass

    i, n = 0, len(lines)
    in_code = False
    while i < n:
        line = lines[i]

        # 代码块
        if line.strip().startswith("```"):
            in_code = not in_code
            if not in_code:
                i += 1
                continue
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i]); i += 1
            in_code = False
            p = doc.add_paragraph()
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"; r.font.size = Pt(9.5)
            p.paragraph_format.left_indent = Inches(0.25)
            i += 1
            continue

        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # 公式块 $$...$$（单行或跨行）；所有分支必须推进 i，防止死循环
        if stripped.startswith("$$"):
            buf = []
            content = stripped[2:]
            if content.endswith("$$") and len(stripped) > 4:
                # 完整单行公式 $$...$$
                buf.append(content[:-2])
                i += 1
            elif content.strip() == "":
                # 独立开界符 $$，跨行直到闭界
                i += 1
                while i < n:
                    t = lines[i].strip()
                    if t == "$$":
                        i += 1
                        break
                    if t.endswith("$$") and len(t) > 2:
                        buf.append(t[:-2]); i += 1
                        break
                    buf.append(lines[i].rstrip()); i += 1
            else:
                # 开界后同行有内容，跨行直到闭界
                buf.append(content.rstrip())
                i += 1
                while i < n:
                    t = lines[i].strip()
                    if t == "$$" or t.endswith("$$"):
                        if t != "$$" and len(t) > 2:
                            buf.append(t[:-2])
                        i += 1
                        break
                    buf.append(lines[i].rstrip()); i += 1
            if buf:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run("$$" + "\n".join(buf) + "$$")
                r.font.name = "Cambria Math"; r.font.size = Pt(11)
                set_cjk(r, "宋体", "Cambria Math")
            continue

        # 标题
        m = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            h = doc.add_heading("", level=level)
            add_runs(h, m.group(2))
            for r in h.runs:
                set_cjk(r, "黑体", "Arial")
            i += 1
            continue

        # 分割线
        if re.match(r"^-{3,}$|^\*{3,}$", stripped):
            doc.add_paragraph("—" * 30).alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # 表格
        if stripped.startswith("|") and i + 1 < n and is_table_sep(lines[i + 1]):
            rows = [split_table_row(stripped)]
            i += 2
            while i < n and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i])); i += 1
            add_table(doc, rows)
            continue

        # 引用
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_runs(p, stripped.lstrip("> "))
            for r in p.runs:
                r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B); r.font.size = Pt(10.5)
            i += 1
            continue

        # 图片
        m = re.match(r"^!\[(.*?)\]\((.+?)\)$", stripped)
        if m:
            img_path = Path(m.group(2))
            if not img_path.is_absolute():
                img_path = base_dir / img_path
            if img_path.exists():
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                try:
                    p.add_run().add_picture(str(img_path), width=Inches(5.5))
                except Exception as e:
                    p.add_run(f"[图片插入失败: {img_path.name} ({e})]")
                if m.group(1):
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = cap.add_run(m.group(1)); r.font.size = Pt(9)
                    r.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            else:
                doc.add_paragraph(f"[图片缺失: {img_path}]")
            i += 1
            continue

        # 列表（两级）
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            indent, marker, content = m.group(1), m.group(2), m.group(3)
            style = "List Number" if marker[0].isdigit() else "List Bullet"
            if len(indent) >= 2:
                style += " 2"
            p = doc.add_paragraph(style=style)
            add_runs(p, content)
            i += 1
            continue

        # 普通段落
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(str(out_path))
    print(f"✓ 已导出 Word 文档: {out_path.resolve()}")


def main():
    parser = argparse.ArgumentParser(description="Markdown 转 Word (.docx)")
    parser.add_argument("input", help="输入的 .md 文件（如 notes.md）")
    parser.add_argument("-o", "--output", default=None, help="输出 .docx 路径（默认: 同名 .docx）")
    args = parser.parse_args()

    md_path = Path(args.input)
    if not md_path.is_file():
        print(f"✗ 找不到输入文件: {md_path}")
        sys.exit(1)
    out_path = Path(args.output) if args.output else md_path.with_suffix(".docx")
    convert(md_path, out_path)


if __name__ == "__main__":
    main()
